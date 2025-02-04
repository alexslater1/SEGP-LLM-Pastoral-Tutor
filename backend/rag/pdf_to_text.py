from io import BytesIO
import PyPDF2
from docx import Document
from pptx import Presentation
import io

async def pdf_to_text(file_contents):
    pdf_file = io.BytesIO(file_contents)
    pdf_reader = PyPDF2.PdfReader(pdf_file)
    text = ""
    for page in pdf_reader.pages:
        text += page.extract_text()
    return text

async def docx_to_text(file_contents):
    doc = Document(BytesIO(file_contents))
    full_text = []
    for paragraph in doc.paragraphs:
        full_text.append(paragraph.text)
    return '\n'.join(full_text)

async def txt_to_text(file_contents):
    return file_contents.decode('utf-8')

async def pptx_to_text(file_contents):
    presentation = Presentation(BytesIO(file_contents))
    text = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
    return '\n'.join(text)
    

async def file_to_text(file_contents, file_name):
    if file_name.endswith('.pdf'):
        return await pdf_to_text(file_contents)
    elif file_name.endswith('.docx'):
        return await docx_to_text(file_contents)
    elif file_name.endswith('.txt'):
        return await txt_to_text(file_contents)
    elif file_name.endswith('.pptx'):
        return await pptx_to_text(file_contents)
    else:
        raise ValueError(f"Unsupported file type")
